# -*- coding: utf-8 -*-
"""v1.18 · el escritor rápido de libros produce lo mismo que pandas, sin fórmulas ni caracteres inválidos."""

import io
import zipfile

import numpy as np
import openpyxl
import pandas as pd
import pytest

from core import excel_safe, fastxlsx
from core import storage as sm
from core.storage import LabStorage


def _read(data: bytes, sheet: str) -> pd.DataFrame:
    return pd.ExcelFile(io.BytesIO(data), engine="openpyxl").parse(sheet, dtype=str, keep_default_na=False).fillna("")


def _fast(frames) -> bytes:
    buffer = io.BytesIO()
    fastxlsx.write_workbook(buffer, frames)
    return buffer.getvalue()


def _pandas(frames) -> bytes:
    buffer = io.BytesIO()
    excel_safe.write_frames(buffer, frames)
    return buffer.getvalue()


FRAME = pd.DataFrame({
    "id": ["a1", "a2", "a3", "a4", "a5"],
    "texto": ["=1+1", "  con espacios  ", "línea 1\nlínea 2", "Ñandú & <b>\"q\"</b>", ""],
    "numero": [1, 2, 3, 4, 5],
    "decimal": [1.5, 0.25, 3.0, 4.125, 5.0],
    "flag": [True, False, True, False, True],
}).astype({"texto": object})


def test_reads_back_the_same_as_the_pandas_path():
    fast, slow = _read(_fast({"hoja": FRAME}), "hoja"), _read(_pandas({"hoja": FRAME}), "hoja")
    pd.testing.assert_frame_equal(fast, slow)
    assert fast["texto"].tolist() == ["=1+1", "  con espacios  ", "línea 1\nlínea 2", "Ñandú & <b>\"q\"</b>", ""]


def test_formulas_stay_text_and_numbers_keep_their_type():
    book = openpyxl.load_workbook(io.BytesIO(_fast({"hoja": FRAME})))
    sheet = book["hoja"]
    assert sheet["B2"].value == "=1+1" and sheet["B2"].data_type == "s"
    assert sheet["C2"].value == 1 and sheet["C2"].data_type == "n"
    assert sheet["D2"].value == 1.5 and sheet["E2"].value is True and sheet["E3"].value is False
    assert sheet["B6"].value is None                                   # el texto vacío es una celda en blanco


def test_invalid_xml_characters_are_removed_and_long_text_is_capped():
    frame = pd.DataFrame({"a": ["pinza\x1d\x00", "x" * 40000, "ok\ufffe"]}).astype(object)
    values = _read(_fast({"h": frame}), "h")["a"].tolist()
    assert values[0] == "pinza" and len(values[1]) == 32767 and values[2] == "ok"


def test_numpy_scalars_nan_and_missing_values():
    frame = pd.DataFrame({"n": np.array([1, 2], dtype=np.int64), "f": np.array([1.5, np.nan]),
                          "b": np.array([True, False])})
    sheet = openpyxl.load_workbook(io.BytesIO(_fast({"h": frame})))["h"]
    assert sheet["A2"].value == 1 and sheet["B2"].value == 1.5 and sheet["B3"].value is None
    assert sheet["C2"].value is True and sheet["C3"].value is False


def test_many_sheets_empty_frames_and_odd_names():
    frames = {"items": pd.DataFrame(columns=["id", "name"]), "trace & events": FRAME, "x" * 40: FRAME}
    book = openpyxl.load_workbook(io.BytesIO(_fast(frames)))
    assert book.sheetnames == ["items", "trace & events", "x" * 31]
    assert [c.value for c in book["items"][1]] == ["id", "name"] and book["items"].max_row == 1


def test_output_is_deterministic_and_a_valid_zip():
    first, second = _fast({"hoja": FRAME}), _fast({"hoja": FRAME})
    assert first == second
    assert zipfile.ZipFile(io.BytesIO(first)).testzip() is None


def test_the_database_round_trips_through_the_fast_writer():
    storage = LabStorage()
    storage.save_item({"name": "=SUMA(1;2)", "item_type": "standalone", "quantity": 7, "description": "a\x1db"},
                      "LAB-FAST-01", is_new=True)
    sm._cached_dfs = None
    item = storage.get_item("LAB-FAST-01")
    assert item["name"] == "=SUMA(1;2)" and item["quantity"] == 7 and item["description"] == "ab"
    sm._validate_workbook_bytes(open(sm.EXCEL_PATH, "rb").read())          # la descarga validada lo acepta


def test_falls_back_to_pandas_if_the_fast_writer_fails(monkeypatch):
    def boom(*args, **kwargs):
        raise RuntimeError("fallo simulado")

    monkeypatch.setattr(fastxlsx, "write_workbook", boom)
    storage = LabStorage()
    storage.save_item({"name": "Lupa", "item_type": "standalone", "quantity": 1}, "LAB-FB-01", is_new=True)
    sm._cached_dfs = None
    assert storage.get_item("LAB-FB-01")["name"] == "Lupa"


@pytest.mark.parametrize("name", ["items", "users", "loans", "trace_events"])
def test_every_database_sheet_keeps_its_columns(name):
    LabStorage()
    frames = sm._read_excel()
    again = _read(_fast(frames), name)
    assert list(again.columns) == sm.SHEET_COLUMNS[name]


def test_unchanged_sheets_are_reused_and_changed_ones_are_rebuilt():
    fastxlsx.clear_cache()
    first = pd.DataFrame({"a": ["x", "y"]}).astype(object)
    other = pd.DataFrame({"b": ["1"]}).astype(object)
    one = _fast({"h1": first, "h2": other})
    cached_xml = fastxlsx._SHEET_CACHE["h2"][1]
    changed = first.copy()
    changed.iloc[1, 0] = "z"
    two = _fast({"h1": changed, "h2": other.copy()})
    assert fastxlsx._SHEET_CACHE["h2"][1] is cached_xml                    # h2 no se volvió a generar
    assert _read(one, "h1")["a"].tolist() == ["x", "y"] and _read(two, "h1")["a"].tolist() == ["x", "z"]
    assert _read(two, "h2")["b"].tolist() == ["1"]


def test_the_workbook_opens_without_openpyxl_style_warnings(recwarn):
    openpyxl.load_workbook(io.BytesIO(_fast({"h": FRAME})))
    assert not [w for w in recwarn if "default style" in str(w.message)]


# --- lectura rápida de encabezados (arranque) ----------------------------------------------------

def test_headers_are_read_from_the_first_row_only(tmp_path):
    LabStorage()
    headers = sm._workbook_headers(sm.EXCEL_PATH)
    assert headers == {sheet: list(columns) for sheet, columns in sm.SHEET_COLUMNS.items()}


def test_headers_fall_back_to_openpyxl_for_files_saved_by_excel(tmp_path):
    path = tmp_path / "excel.xlsx"
    pd.DataFrame({"id": ["1"], "name & más": ["x"]}).to_excel(path, sheet_name="items", index=False)   # textos compartidos
    assert sm._workbook_headers(str(path)) == {"items": ["id", "name & más"]}


def test_headers_of_an_empty_sheet_and_special_names(tmp_path):
    path = tmp_path / "raro.xlsx"
    fastxlsx.write_workbook(str(path), {"a&b": pd.DataFrame(columns=["x<y", "z"]), "vacía": pd.DataFrame()})
    assert sm._workbook_headers(str(path))["a&b"] == ["x<y", "z"]


def test_old_workbooks_missing_columns_are_completed_on_startup():
    storage = LabStorage()
    frames = sm._read_excel()
    frames["users"] = frames["users"].drop(columns=["session_epoch"])
    buffer = io.BytesIO()
    fastxlsx.write_workbook(buffer, frames)
    with open(sm.EXCEL_PATH, "wb") as handle:
        handle.write(buffer.getvalue())
    sm._cached_dfs = None
    LabStorage()                                              # detecta la columna que falta y migra
    assert "session_epoch" in sm._workbook_headers(sm.EXCEL_PATH)["users"]
    assert storage.get_all_users() == []
