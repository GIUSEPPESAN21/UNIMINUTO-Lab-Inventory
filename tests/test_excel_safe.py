# -*- coding: utf-8 -*-
"""v1.16 · texto de personas en el Excel: sin fórmulas accidentales ni caracteres de control."""

import io

import openpyxl
import pandas as pd
import pytest

from core import excel_safe, ludica, reports
from core import storage as sm
from core.storage import LabStorage

FORMULA = '=HYPERLINK("http://ejemplo.invalid","clic")'


def _reload(storage):
    sm._cached_dfs = None            # fuerza leer el archivo, no la caché
    return storage


def _item(name, **extra):
    return {"name": name, "item_type": "standalone", "quantity": 1, "location": "Estante", **extra}


# --- el dato sobrevive al ciclo guardar → recargar -----------------------------------------

@pytest.mark.parametrize("text", ["=1+1", FORMULA, "=", "+57 300", "-x", "@usuario"])
def test_text_that_looks_like_a_formula_survives_a_reload(text):
    storage = LabStorage()
    storage.save_item(_item(text, description=text), "LAB-F-01", is_new=True)
    item = _reload(storage).get_item("LAB-F-01")
    assert item["name"] == text and item["description"] == text


def test_formula_text_is_stored_as_text_cell_in_the_file():
    storage = LabStorage()
    storage.save_item(_item("=1+1"), "LAB-F-02", is_new=True)
    sheet = openpyxl.load_workbook(sm.EXCEL_PATH)["items"]
    cells = [cell for row in sheet.iter_rows() for cell in row if cell.value == "=1+1"]
    assert cells and all(cell.data_type == "s" for cell in cells)    # nunca 'f' (fórmula)


def test_control_characters_are_removed_instead_of_breaking_the_save():
    storage = LabStorage()
    # \x1d (separador de grupo) es lo que algunas pistolas de códigos de barras añaden al final
    storage.save_item(_item("Pinza\x1d", description="nota\x01 con\x1f control"), "LAB-F-03", is_new=True)
    item = _reload(storage).get_item("LAB-F-03")
    assert item["name"] == "Pinza" and item["description"] == "nota con control"


def test_control_characters_do_not_poison_the_deferred_flush(monkeypatch):
    monkeypatch.setattr(sm, "safe_secret", lambda key, default=None: "deferred" if key == "SYNC_MODE" else default)
    storage = LabStorage()
    storage.save_item(_item("Lupa\x1d"), "LAB-F-04", is_new=True)      # solo cambia la caché
    sm._persist_file_if_dirty()                                         # el vaciado al archivo ya no lanza
    assert sm._sync_state["file_dirty"] is False                        # y no queda reintentando para siempre
    assert _reload(storage).get_item("LAB-F-04")["name"] == "Lupa"


def test_tabs_and_newlines_are_kept():
    storage = LabStorage()
    storage.save_item(_item("A", description="línea 1\nlínea 2\tcolumna"), "LAB-F-05", is_new=True)
    assert _reload(storage).get_item("LAB-F-05")["description"] == "línea 1\nlínea 2\tcolumna"


# --- el archivo descargado por la persona no ejecuta fórmulas -----------------------------------

def test_downloadable_export_keeps_formula_text_as_text():
    items = [{"id": "LAB-1", "name": FORMULA, "quantity": 1}]
    users = [{"id": "u1", "full_name": "=cmd|' /C calc'!A0", "password_hash": "x", "role": "estudiante"}]
    book = openpyxl.load_workbook(reports.export_full_database(items, users, []))
    for sheet, expected in (("items", FORMULA), ("users", "=cmd|' /C calc'!A0")):
        cells = [cell for row in book[sheet].iter_rows() for cell in row if cell.value == expected]
        assert cells and all(cell.data_type == "s" for cell in cells)
    assert "password_hash" not in [cell.value for cell in book["users"][1]]


def test_write_frames_handles_pandas_string_dtype_and_numbers():
    frame = pd.DataFrame({"a": pd.Series(["=x", "ok"], dtype="string"), "b": [1, 2]})
    buffer = io.BytesIO()
    excel_safe.write_frames(buffer, {"hoja": frame})
    back = pd.read_excel(io.BytesIO(buffer.getvalue()), dtype=str, keep_default_na=False)
    assert back["a"].tolist() == ["=x", "ok"] and back["b"].tolist() == ["1", "2"]


# --- CSV: Excel evalúa lo que empieza por = + - @ ---------------------------------------------

@pytest.mark.parametrize("value, expected", [
    ("=1+1", "'=1+1"), ("@SUMA(A1)", "'@SUMA(A1)"), ("+57 300 000", "'+57 300 000"),
    ("-cmd", "'-cmd"), ("\tcmd", "'\tcmd"),
    ("-5", "-5"), ("+3.5", "+3.5"), ("LAB-001", "LAB-001"), ("", ""), (None, None), (7, 7),
])
def test_csv_safe(value, expected):
    assert excel_safe.csv_safe(value) == expected


def test_lúdica_csv_neutralizes_a_code_that_starts_like_a_formula(monkeypatch):
    row = dict.fromkeys(ludica.COLUMNS, "x")
    row["item_codigo"] = "=1+1"
    monkeypatch.setattr(ludica, "event_rows", lambda *args, **kwargs: [row])
    admin = {"id": "a1", "role": "maestro"}
    lines = ludica.events_csv(LabStorage(), admin).decode("utf-8").splitlines()
    assert lines[0] == ",".join(ludica.COLUMNS)
    assert "'=1+1" in lines[1] and ",=1+1," not in lines[1]
